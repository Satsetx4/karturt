import json
import os
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from unittest.mock import patch

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

import scripts.validate_sheet2_workbook as sheet2_validator
from scripts.validate_sheet2_workbook import WorkbookValidationError, read_and_validate, write_private_map


def make_synthetic_workbook(path: Path, partial_range: bool = False) -> None:
    workbook = openpyxl.Workbook()
    instructions = workbook.active
    instructions.title = "Sheet1"
    instructions["A1"] = "Untrusted worksheet content must not affect this importer."

    sheet = workbook.create_sheet("Sheet2")
    headers = {
        "B2": "Nomor Rumah",
        "C2": "Nama Kepala Rumah Tangga",
        "D2": "WhatsApp",
        "F2": "Tarif Bulanan",
        "G2": "Riwayat Pembayaran",
        "H2": "Riwayat Pembayaran",
        "I2": "Username",
        "J2": "PIN",
        "K2": "Akses",
        "G3": "Awal Bayar",
        "H3": "Terakhir Bayar",
    }
    for cell, value in headers.items():
        sheet[cell] = value

    end_months = [12] * 16 + [11] * 2 + [9] * 36 + [8] * 11 + [None] * 8
    roles = ["Ketua RT", "Bendahara", "Admin Sistem"] + ["Warga"] * 70
    for index in range(73):
        row = index + 4
        sheet.cell(row, 2, f"SYNTH-HOUSE-{index + 1:03}")
        sheet.cell(row, 3, f"Synthetic Resident {index + 1:03}")
        sheet.cell(row, 4, "0000000000")
        sheet.cell(row, 6, 1000 if index < 57 else 2000)
        end_month = end_months[index]
        if end_month is not None:
            sheet.cell(row, 7, "Januari")
            sheet.cell(row, 8, ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"][end_month - 1])
        sheet.cell(row, 9, f"SYNTH{index + 1:03}")
        sheet.cell(row, 10, "123456")
        sheet.cell(row, 11, roles[index])
    if partial_range:
        sheet.cell(4, 8).value = None
    workbook.save(path)
    workbook.close()


class Sheet2WorkbookValidatorTests(unittest.TestCase):
    def test_synthetic_roster_aggregates_and_sheet1_is_ignored(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.xlsx"
            make_synthetic_workbook(path)
            summary, private_map = read_and_validate(path)

        self.assertEqual(summary["status"], "PASS")
        self.assertFalse(summary["sheet1_read"])
        self.assertEqual(summary["resident_rows"], 73)
        self.assertEqual(summary["tariff_class_counts"], [57, 16])
        self.assertEqual(summary["paid_household_months"], 626)
        self.assertEqual(summary["prepaid_future_household_months"], {11: 18, 12: 16})
        self.assertEqual(summary["overdue_household_months"], 138)
        self.assertEqual(summary["unpaid_not_yet_due_household_months"], 112)
        self.assertEqual(summary["expected_total_accounts"], 76)
        self.assertEqual(summary["new_accounts_to_create"], 75)
        self.assertEqual(summary["reused_system_admin_account_count"], 1)
        self.assertTrue(all("pin" not in entry for entry in private_map["entries"]))
        self.assertNotIn("Synthetic Resident", json.dumps(summary))
        self.assertNotIn("SYNTH-HOUSE", json.dumps(summary))

    def test_partial_range_fails_closed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.xlsx"
            make_synthetic_workbook(path, partial_range=True)
            summary, _ = read_and_validate(path)

        self.assertEqual(summary["status"], "BLOCKED")
        self.assertEqual(summary["partial_payment_ranges"], 1)
        self.assertEqual(summary["validation_errors"]["partial_payment_range"], 1)

    def test_pin_payload_is_opt_in_and_cannot_be_persisted_as_private_map(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "synthetic.xlsx"
            make_synthetic_workbook(path)
            _, ordinary_map = read_and_validate(path)
            _, bridge_map = read_and_validate(path, include_pins=True)

            self.assertTrue(all("pin" not in entry for entry in ordinary_map["entries"]))
            self.assertTrue(all(entry["pin"] == "123456" for entry in bridge_map["entries"]))
            with self.assertRaises(WorkbookValidationError):
                write_private_map(root, Path("data/private/synthetic.json"), bridge_map)

    def test_private_bridge_requires_importer_parent_and_emits_minimal_sheet2_payload(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.xlsx"
            make_synthetic_workbook(path)
            batch_id = "4e87b570-f17f-4e0b-b01a-0f2a06f01185"
            bridge_stdout = StringIO()
            with patch.dict(os.environ, {
                "KARTURT_SHEET2_IMPORT_BRIDGE": "1",
                "KARTURT_SHEET2_IMPORT_PARENT_PID": "12345",
                "KARTURT_SHEET2_IMPORT_BRIDGE_NONCE": "a" * 32,
                "KARTURT_SHEET2_IMPORT_BATCH_ID": batch_id,
            }), patch.object(sheet2_validator.os, "getppid", return_value=12345), \
                 patch.object(sys, "argv", ["validate_sheet2_workbook.py", "--_emit-private-import-payload", str(path)]), \
                 redirect_stdout(bridge_stdout):
                status = sheet2_validator.main()

            payload = json.loads(bridge_stdout.getvalue())
            self.assertEqual(status, 0)
            self.assertEqual(payload["importBatchId"], batch_id)
            self.assertEqual(payload["sourceSheet"], "Sheet2")
            self.assertEqual(len(payload["entries"]), 73)
            self.assertTrue(all(entry["pin"] == "123456" for entry in payload["entries"]))
            self.assertTrue(all("phone" not in entry and "sourceRow" not in entry for entry in payload["entries"]))
            self.assertNotIn("Sheet1", bridge_stdout.getvalue())

            denied_stdout = StringIO()
            with patch.dict(os.environ, {"KARTURT_SHEET2_IMPORT_BRIDGE": "0"}, clear=True), \
                 patch.object(sys, "argv", ["validate_sheet2_workbook.py", "--_emit-private-import-payload", str(path)]), \
                 redirect_stdout(denied_stdout):
                denied_status = sheet2_validator.main()
            self.assertEqual(denied_status, 2)
            self.assertEqual(denied_stdout.getvalue(), "")


if __name__ == "__main__":
    unittest.main()
